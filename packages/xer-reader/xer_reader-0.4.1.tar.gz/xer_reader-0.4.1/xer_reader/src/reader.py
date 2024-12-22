"""
This module contains the `XerReader class, which represents the contents of
an XER file exported from Primavera P6.

"""

import csv
import json
import re
from datetime import datetime
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from xer_reader.src.table import XerTable, UnrecognizedTable
from xer_reader.src.table_data import table_data

DATE_FORMAT = "%Y-%m-%d"
REQUIRED_TABLES = {"CALENDAR", "CURRTYPE", "PROJECT", "PROJWBS"}


class XerReader:
    """Open an XER file exported from Primavera P6 and read its contents."""

    CODEC = "cp1252"
    file_name: str
    """XER file name"""
    data: str
    """XER file raw data as tab seperated text"""

    def __init__(self, file: str | Path | BinaryIO) -> None:
        self.file_name, self.data = _read_file(file)

        _file_info = _parse_file_info(self.data)
        self.currency: str = _file_info[7]
        """(str) Currency type set in P6"""
        self.export_version: str = _file_info[0]
        """(str) P6 Version"""
        self.export_date: datetime = datetime.strptime(_file_info[1], DATE_FORMAT)
        """(datetime) Date the XER file was exported"""
        self.export_user: str = _file_info[4]
        """(str) P6 user name that exported the XER file"""

    def check_errors(self) -> list[str]:
        """Check XER file for missing tables and orphan data

        Returns:
            list[str]: Descriptions of missing information
        """
        errors = set()

        id_map = {
            data["key"]: table for table, data in table_data.items() if data["key"]
        }

        tables = self.to_dict()

        # Check for minimum tables required to be in the XER
        for name in REQUIRED_TABLES:
            if name not in tables:
                errors.add(f"Missing Required Table {name}")

        # Check for required table pairs
        for table in tables.values():
            for table2 in table.depends:
                if table2 not in tables:
                    errors.add(f"Missing Table {table2} Required for Table {table}")

            for row in table.entries():
                for key, val in row.items():
                    if val == "":
                        continue
                    if not key.endswith("_id"):
                        continue
                    if key == "parent_wbs_id" and row["proj_node_flag"] == "Y":
                        continue
                    clean_key = key if key in id_map else _clean_foreign_key_label(key)
                    if clean_key:
                        if check_table := tables.get(id_map[clean_key]):
                            for entry in check_table.entries():
                                if entry.get(clean_key, "") == val:
                                    break
                            else:
                                errors.add(
                                    f"Orphan data {key} [{val}] in table {table}"
                                )

        return list(errors)

    def delete_tables(self, *table_names: str) -> str:
        """
        Delete tables from XER file.
        Does not modify `XerReader.data` attribute, but returns a new string.

        Args:
            *table_names (str): table names to remove from XER file

        Returns:
            str: XER File String with tables removed
        """
        if not table_names:
            raise ValueError("Must pass at least one table name")

        rev_data = self.data
        for name in table_names:
            table_search = re.compile(rf"%T\t{name.upper()}\r?\n(.|\s)*?(?=%T|%E)")
            rev_data = table_search.sub("", rev_data)
        return rev_data

    def get_table_names(self) -> list[str]:
        """Get list of table names included in the XER file.

        Returns:
            list[str]: list of table names
        """
        table_names = re.compile(r"(?<=%T\t)[A-Z]+")
        return table_names.findall(self.data)

    def get_table_str(self, table_name: str) -> str:
        """Get string for a specific table in the XER file.

        Args:
            table_name (str): Name of table

        Returns:
            str: Table header and rows
        """
        re_search = re.compile(rf"(?<=%T\t{table_name.upper()}\n)(.|\s)*?(?=%T|%E)")
        if found_table := re_search.search(self.data):
            return re.sub(r"%[TFR]\t", "", found_table.group())
        return ""

    def has_table(self, table_name: str) -> bool:
        """Check if a table is included in the XER file.

        Args:
            table_name (str): table name

        Returns:
            bool: True if found; False if not found
        """
        return f"%T\t{table_name.upper()}" in self.data

    def to_dict(self) -> dict[str, XerTable]:
        """
        Parse tables into a dictionary with the table name as the key
        and a `Table` object as the value.

        Returns:
            dict[str, Table]: dict of XER Tables
        """
        tables = {}
        for table_str in self.data.split("%T\t")[1:]:
            try:
                table = XerTable(table_str)
                tables[table.name] = table
            except UnrecognizedTable:
                continue
        return tables

    def to_csv(
        self,
        file_directory: str | Path = Path.cwd(),
        table_names: list[str] = [],
        delimeter: str = "\t",
    ) -> None:
        """
        Generate a CSV file for each table in the XER file.
        Uses `tab` as the delimiter by default.

        Args:
            file_directory (str | Path, optional): Directory to save CSV files. [Defaults to current working directory]
            table_names (list, optional): List of table names to save to CSV files. If empty, all tables will be saved.
            delimeter (str, optional): CSV delimeter. [Default is a `tab`]
        """
        names = [name.upper() for name in table_names]
        for table in self.to_dict().values():
            if not table_names or table.name in names:
                _write_table_to_csv(
                    f"{self.file_name}_{table.name}",
                    table,
                    Path(file_directory),
                    delimeter,
                )

    def to_excel(self, file_directory: str | Path = Path.cwd()) -> None:
        """
        Generate an Excel file with each table in the XER file on a seperate worksheet.

        """
        wb = Workbook()
        ws = wb.active
        ws.title = "ERMHDR"
        ws.append(_parse_file_info(self.data))

        for name, table in self.to_dict().items():
            new_ws = wb.create_sheet(name)
            new_ws.append(table.labels)
            for entry in table.rows:
                new_ws.append(entry)

            tab = Table(displayName=name, ref=new_ws.calculate_dimension())

            new_ws.add_table(tab)

        wb.save(Path(file_directory, f"{self.file_name}.xlsx"))

    def to_json(self, *tables: str) -> str:
        """Generate a json compliant string representation of tables in the XER file

        Returns:
            str: json compliant string representation of XER tables
        """
        out_data = {}
        if not tables:
            out_data = {
                name: _entry_by_key(table) for name, table in self.to_dict().items()
            }
        else:
            out_data = {
                name: _entry_by_key(table)
                for name, table in self.to_dict().items()
                if name in tables
            }
        json_data = {self.file_name: {**out_data}}
        return json.dumps(json_data, indent=2)


def _clean_foreign_key_label(label: str) -> str | None:
    prefixes = ("base_", "last_", "new_", "parent_", "pred_")
    for prefix in prefixes:
        if label.startswith(prefix):
            return label.replace(prefix, "")
    return


def _entry_by_key(table: XerTable) -> dict | list:
    if not table.key:
        return table.entries(serialize=True)
    return {entry[table.key]: entry for entry in table.entries(serialize=True)}


def _parse_file_info(data: str) -> list[str]:
    """Parse file header"""
    ermhdr = re.search(r"(?<=ERMHDR\t).+", data)
    if not ermhdr:
        raise ValueError("Invalid XER File")
    return ermhdr.group().split("\t")


def _read_file(file: str | Path | BinaryIO) -> tuple[str, str]:
    file_contents = ""
    file_name = ""
    if isinstance(file, (str, Path)):
        # Path directory to file
        file_name = Path(file).stem
        with open(file, encoding=XerReader.CODEC, errors="ignore") as f:
            file_contents = f.read()
    else:
        # Binary file from requests, Flask, FastAPI, etc...
        file_contents = file.read().decode(XerReader.CODEC, errors="ignore")
        file_name = file.name

    if not file_contents.startswith("ERMHDR"):
        raise ValueError(f"ValueError: {file_name} is invalid XER file")

    return file_name, file_contents


def _write_table_to_csv(
    name: str, table: XerTable, file_directory: Path, delimeter
) -> None:
    with file_directory.joinpath(f"{name}.csv").open("w") as f:
        writer = csv.writer(f, delimiter=delimeter)
        writer.writerow(table.labels)
        for row in table.rows:
            writer.writerow(row)
    f.close()
